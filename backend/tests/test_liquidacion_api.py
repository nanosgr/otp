"""API de liquidación: calcular, cerrar/reabrir, resumen y casos de error (SQLite + semilla real de reglas)."""
from datetime import date
from decimal import Decimal

import pytest
from fastapi.testclient import TestClient

from app.core import rbac
from app.core.security import create_access_token, get_password_hash
from app.db.database import get_db
from app.db.init_prevision import init_prevision
from app.db.seed_reglas import seed_reglas
from app.main import app
from app.models import prevision as m
from app.models.models import Role, User
from app.services import motor

pytestmark = pytest.mark.skipif(motor._engine is None, reason="otp_engine no instalado")


@pytest.fixture(autouse=True)
def _clear_cache():
    rbac.invalidate_policy_cache()
    yield
    rbac.invalidate_policy_cache()


@pytest.fixture
def client(db):
    seed_reglas(db)
    init_prevision(db)
    app.dependency_overrides[get_db] = lambda: db
    with TestClient(app) as c:
        yield c
    app.dependency_overrides.clear()


def _headers(db, name, role):
    r = db.query(Role).filter(Role.name == role).first()
    u = User(username=name, email=f"{name}@e.com", hashed_password=get_password_hash("x"))
    u.roles.append(r)
    db.add(u)
    db.commit()
    db.refresh(u)
    return {"Authorization": "Bearer " + create_access_token({"sub": u.username, "token_version": u.token_version})}


@pytest.fixture
def admin(db, client):
    return _headers(db, "adm", "Administrador")


@pytest.fixture
def liquidador(db, client):
    return _headers(db, "liq", "Liquidador")


def _causante(db, tipo="superior", cargo=None, cargos=1):
    c = m.Causante(dni="20111222", apellido="Pérez", nombre="Juan", tipo_personal=tipo)
    db.add(c)
    db.commit()
    db.refresh(c)
    for i in range(cargos):
        db.add(m.CargoSecuencia(causante_id=c.id, secuencia=i + 1, **(cargo or {"clase": 16, "anios_antiguedad": 30, "porcentaje_retiro": 100})))
    db.commit()
    return c


def _liq(db, causante, **kw):
    datos = dict(periodo="2025-03", tipo="retiro", fecha_desde=date(2024, 7, 1), fecha_hasta=date(2025, 3, 31))
    datos.update(kw)
    liq = m.Liquidacion(causante_id=causante.id, **datos)
    db.add(liq)
    db.commit()
    db.refresh(liq)
    return liq


def test_calcular_retiro_persiste_recibo_tramos_y_totales(db, client, admin):
    liq = _liq(db, _causante(db))
    r = client.post(f"/api/v1/liquidaciones/{liq.id}/calcular", headers=admin)
    assert r.status_code == 200, r.text
    res = r.json()
    rec = res["recibos"][0]
    tramos = rec["tramos"]
    assert [t["es_sac"] for t in tramos].count(True) == 1                 # SAC 2do semestre 2024
    assert tramos[0]["desde"] == "2024-07-01" and tramos[-1]["hasta"] == "2025-03-31"
    # crédito = suma de tramos; líquido = crédito - descuentos (14% retiro: 8 + 5 + 0,25 + 0,75)
    credito = sum(Decimal(t["importe"]) for t in tramos)
    assert abs(Decimal(rec["credito"]) - credito) <= Decimal("0.005")
    assert abs(Decimal(rec["total_descuentos"]) - Decimal(rec["credito"]) * Decimal("0.14")) < Decimal("0.02")
    assert Decimal(rec["liquido"]) == Decimal(rec["credito"]) - Decimal(rec["total_descuentos"])
    assert {d["codigo"] for d in rec["descuentos"]} == {"DESC_LEY_FEDERAL", "DESC_OSEP_CUOTA", "DESC_OSEP_DIRECTO", "DESC_OSEP_INCAPACIDAD", "ANTICIPO"}
    assert res["liquidacion"]["total_liquido"] == rec["liquido"]
    assert res["liquidacion"]["calculada_at"]

    # recalcular reemplaza (no duplica)
    r2 = client.post(f"/api/v1/liquidaciones/{liq.id}/calcular", headers=admin)
    assert len(r2.json()["recibos"]) == 1
    assert len(client.get(f"/api/v1/tramos-retroactivos/?liquidacion_id={liq.id}&size=100", headers=admin).json()["items"]) == len(tramos)


def test_anticipo_se_descuenta_del_liquido(db, client, admin):
    liq = _liq(db, _causante(db), anticipo_importe=Decimal("1000"))
    base = client.post(f"/api/v1/liquidaciones/{liq.id}/calcular", headers=admin).json()["recibos"][0]
    assert {d["codigo"]: d["importe"] for d in base["descuentos"]}["ANTICIPO"] == "1000.0000"
    assert Decimal(base["liquido"]) == Decimal(base["credito"]) * Decimal("0.86") - Decimal("1000") or abs(
        Decimal(base["liquido"]) - (Decimal(base["credito"]) * Decimal("0.86") - Decimal("1000"))) < Decimal("0.02")


def test_pension_un_recibo_por_beneficiario_con_art37(db, client, admin):
    c = _causante(db)
    db.add_all([
        m.Beneficiario(causante_id=c.id, apellido="P", nombre="Ana", parentesco="conyuge", porcentaje=Decimal(60)),
        m.Beneficiario(causante_id=c.id, apellido="P", nombre="Leo", parentesco="hijo", porcentaje=Decimal(40), art37=True),
    ])
    db.commit()
    liq = _liq(db, c, tipo="pension")
    res = client.post(f"/api/v1/liquidaciones/{liq.id}/calcular", headers=admin).json()
    r1, r2 = res["recibos"]
    assert r1["beneficiario"]["nombre"] == "Ana" and r2["beneficiario"]["nombre"] == "Leo"
    # descuentos de pensión: solo OSEP (6%), sin Ley Federal
    assert {d["codigo"] for d in r1["descuentos"]} == {"DESC_OSEP_CUOTA", "DESC_OSEP_DIRECTO", "DESC_OSEP_INCAPACIDAD", "ANTICIPO"}
    # el haber del hijo (40% + art. 37) supera proporcionalmente al del cónyuge (60%) por el 5% extra
    h1, h2 = Decimal(r1["tramos"][0]["haber_mensual"]), Decimal(r2["tramos"][0]["haber_mensual"])
    assert h2 > h1 * Decimal(40) / Decimal(60)
    assert Decimal(res["liquidacion"]["total_liquido"]) == Decimal(r1["liquido"]) + Decimal(r2["liquido"])


def test_beneficiario_con_alta_y_baja_dentro_del_rango(db, client, admin):
    c = _causante(db)
    db.add(m.Beneficiario(causante_id=c.id, apellido="P", nombre="Ana", porcentaje=Decimal(100),
                          fecha_alta=date(2024, 10, 16), fecha_baja=date(2025, 1, 15)))
    db.commit()
    liq = _liq(db, c, tipo="pension")
    rec = client.post(f"/api/v1/liquidaciones/{liq.id}/calcular", headers=admin).json()["recibos"][0]
    assert rec["rango"] == {"desde": "2024-10-16", "hasta": "2025-01-15"}
    meses = sum(Decimal(t["meses"]) for t in rec["tramos"])
    assert meses == Decimal(3)                       # 16/10 .. 15/01 = 3 meses comerciales
    # el semestre cierra el 31/12 dentro del rango: SAC proporcional a los 2,5 meses trabajados en él
    sac = [t for t in rec["tramos"] if t["es_sac"]]
    assert len(sac) == 1
    assert abs(Decimal(sac[0]["importe"]) - Decimal(sac[0]["haber_mensual"]) / 2 * Decimal("2.5") / 6) < Decimal("0.001")


def test_porcentaje_de_retiro_sale_de_la_tabla_segun_tipo_de_personal(db, client, admin):
    # sin porcentaje manual: superior usa la tabla de 30 años; subalterno la de 25
    def haber(tipo, anios):
        db.query(m.Liquidacion).delete()
        db.query(m.CargoSecuencia).delete()
        db.query(m.Causante).delete()
        db.commit()
        c = _causante(db, tipo=tipo, cargo={"clase": 12, "anios_antiguedad": anios})
        liq = _liq(db, c, fecha_desde=date(2025, 3, 1), fecha_hasta=date(2025, 3, 31))
        rec = client.post(f"/api/v1/liquidaciones/{liq.id}/calcular", headers=admin).json()["recibos"][0]
        total = sum(Decimal(x["importe"]) for x in rec["conceptos_haber"] if x["columna"] == "REMUNERATIVO")
        retiro = next(Decimal(x["importe"]) for x in rec["conceptos_haber"] if x["codigo"] == "HABER_RETIRO")
        return retiro / total * 100
    assert abs(haber("superior", 27) - Decimal(91)) < Decimal("0.01")      # tabla N: 27 años = 91%
    assert abs(haber("subalterno", 27) - Decimal(100)) < Decimal("0.01")   # tabla K: 100% desde los 25
    assert abs(haber("superior", 20) - Decimal(65)) < Decimal("0.01")
    assert haber("superior", 9) == 0                                        # menos de 10 años: sin retiro


def test_cerrar_reabrir_e_inmutabilidad(db, client, admin, liquidador):
    liq = _liq(db, _causante(db))
    assert client.post(f"/api/v1/liquidaciones/{liq.id}/cerrar", headers=liquidador).status_code == 409   # sin calcular
    assert client.post(f"/api/v1/liquidaciones/{liq.id}/calcular", headers=liquidador).status_code == 200
    assert client.post(f"/api/v1/liquidaciones/{liq.id}/cerrar", headers=liquidador).json()["liquidacion"]["estado"] == "CERRADA"
    assert client.post(f"/api/v1/liquidaciones/{liq.id}/calcular", headers=liquidador).status_code == 409  # cerrada
    # el liquidador no puede reabrir; el administrador sí
    assert client.post(f"/api/v1/liquidaciones/{liq.id}/reabrir", headers=liquidador).status_code == 403
    assert client.post(f"/api/v1/liquidaciones/{liq.id}/reabrir", headers=admin).json()["liquidacion"]["estado"] == "ABIERTA"
    assert client.post(f"/api/v1/liquidaciones/{liq.id}/calcular", headers=liquidador).status_code == 200


@pytest.mark.parametrize("preparar, codigo, texto", [
    (lambda db: (lambda c: (db.query(m.CargoSecuencia).delete(), db.commit(), c)[-1])(_causante(db)), 422, "no tiene un cargo"),
])
def test_errores_de_datos(db, client, admin, preparar, codigo, texto):
    liq = _liq(db, preparar(db))
    r = client.post(f"/api/v1/liquidaciones/{liq.id}/calcular", headers=admin)
    assert r.status_code == codigo and texto in r.text


def test_errores_de_negocio(db, client, admin):
    c = _causante(db)
    # antes de la primera vigencia cargada (2022-05-01)
    liq = _liq(db, c, fecha_desde=date(2021, 1, 1), fecha_hasta=date(2021, 3, 31))
    r = client.post(f"/api/v1/liquidaciones/{liq.id}/calcular", headers=admin)
    assert r.status_code == 422 and "reglas" in r.text and "2022-05-01" in r.text
    # reajuste no implementado
    liq2 = _liq(db, c, tipo="reajuste")
    assert client.post(f"/api/v1/liquidaciones/{liq2.id}/calcular", headers=admin).status_code == 501
    # pensión sin beneficiarios / porcentajes > 100
    liq3 = _liq(db, c, tipo="pension")
    assert "al menos un beneficiario" in client.post(f"/api/v1/liquidaciones/{liq3.id}/calcular", headers=admin).text
    db.add_all([m.Beneficiario(causante_id=c.id, apellido="A", nombre="A", porcentaje=Decimal(70)),
                m.Beneficiario(causante_id=c.id, apellido="B", nombre="B", porcentaje=Decimal(50))])
    db.commit()
    r = client.post(f"/api/v1/liquidaciones/{liq3.id}/calcular", headers=admin)
    assert r.status_code == 422 and "120" in r.text
    # rango inválido
    liq4 = _liq(db, c, fecha_desde=date(2025, 3, 31), fecha_hasta=date(2025, 3, 1))
    assert client.post(f"/api/v1/liquidaciones/{liq4.id}/calcular", headers=admin).status_code == 422


def test_clase_inexistente_en_la_escala_es_error_de_calculo(db, client, admin):
    c = _causante(db, cargo={"clase": 19, "anios_antiguedad": 30, "porcentaje_retiro": 100})   # la escala no tiene clase 19
    liq = _liq(db, c)
    r = client.post(f"/api/v1/liquidaciones/{liq.id}/calcular", headers=admin)
    assert r.status_code == 422 and "concepto 10" in r.text and "ESCALA_CLASE" in r.text


def test_reglas_sembradas_son_validas_para_el_motor(db, client):
    from app.api.prevision import reglas_desde_db
    problemas = motor.validar_reglas(reglas_desde_db(db))
    assert [p for p in problemas if p["nivel"] == "error"] == []


def test_planillas_xlsx_y_pdf(db, client, admin):
    import io
    import openpyxl
    c = _causante(db)
    c.apellido = "=cmd|' /C calc'!A0"          # un apellido con forma de fórmula no debe ejecutarse
    db.add(c)
    db.add(m.Beneficiario(causante_id=c.id, apellido="+Ana", nombre="Ana", porcentaje=Decimal(100)))
    db.commit()
    liq = _liq(db, c, tipo="pension")
    assert client.get(f"/api/v1/liquidaciones/{liq.id}/planilla.pdf", headers=admin).status_code == 409   # sin calcular
    res = client.post(f"/api/v1/liquidaciones/{liq.id}/calcular", headers=admin).json()

    x = client.get(f"/api/v1/liquidaciones/{liq.id}/planilla.xlsx", headers=admin)
    assert x.status_code == 200 and "attachment" in x.headers["content-disposition"]
    wb = openpyxl.load_workbook(io.BytesIO(x.content))
    ws = wb["Recibo 1"]
    assert ws["B4"].data_type == "s" and ws["B4"].value.startswith("=cmd")      # texto, no fórmula
    valores = {ws[f"A{r}"].value: ws[f"E{r}"].value for r in range(13, ws.max_row + 1) if ws[f"A{r}"].value}
    assert abs(valores["Total crédito líquido"] - float(res["recibos"][0]["liquido"])) < 0.005
    assert "Haber 1" in wb.sheetnames

    p = client.get(f"/api/v1/liquidaciones/{liq.id}/planilla.pdf", headers=admin)
    assert p.status_code == 200 and p.content.startswith(b"%PDF") and p.headers["content-type"] == "application/pdf"


# ---------------------------------------------------------------- dos cargos (secuencias)

def _cargos(db, causante, *cargos):
    db.query(m.CargoSecuencia).filter(m.CargoSecuencia.causante_id == causante.id).delete()
    for i, c in enumerate(cargos, start=1):
        db.add(m.CargoSecuencia(causante_id=causante.id, secuencia=c.pop("secuencia", i), **c))
    db.commit()


def _haber_mensual(client, admin, db, causante, **liq_kw):
    liq = _liq(db, causante, fecha_desde=date(2025, 3, 1), fecha_hasta=date(2025, 3, 31), **liq_kw)
    rec = client.post(f"/api/v1/liquidaciones/{liq.id}/calcular", headers=admin).json()["recibos"][0]
    return Decimal(rec["tramos"][0]["haber_mensual"]), rec


A = {"clase": 16, "anios_antiguedad": 30, "titulo": "grado", "porcentaje_retiro": 100}
B = {"clase": 9, "anios_antiguedad": 10, "titulo_pregrado_nivel": 2}


def test_dos_cargos_el_haber_es_el_promedio_ponderado_de_las_secuencias(db, client, admin):
    c = _causante(db)
    _cargos(db, c, dict(A))
    solo_a, _ = _haber_mensual(client, admin, db, c)
    _cargos(db, c, dict(B, porcentaje_retiro=100))
    solo_b, _ = _haber_mensual(client, admin, db, c)
    _cargos(db, c, dict(A, porcentaje_secuencia=Decimal(60)), dict(B, porcentaje_secuencia=Decimal(40)))
    mixto, rec = _haber_mensual(client, admin, db, c)
    assert abs(mixto - (solo_a * Decimal("0.6") + solo_b * Decimal("0.4"))) < Decimal("0.001")
    # el resumen expone el aporte de cada secuencia y el detalle etiquetado por secuencia
    assert [(x["secuencia"], x["porcentaje"]) for x in rec["cargos"]] == [(1, "60.0000"), (2, "40.0000")]
    assert {x["secuencia"] for x in rec["conceptos_haber"]} == {1, 2, None}
    assert sum(1 for x in rec["conceptos_haber"] if x["codigo"] == "10") == 2      # la clase, una vez por secuencia
    ponderado = sum(Decimal(x["total"]) * Decimal(x["porcentaje"]) / 100 for x in rec["cargos"])
    assert abs(Decimal(rec["haber_ponderado"]) - ponderado) < Decimal("0.0001")


def test_el_porcentaje_de_retiro_sale_del_cargo_principal(db, client, admin):
    c = _causante(db)
    # se cargan en orden inverso: la secuencia 1 (con 100%) manda aunque la 2 tenga otro valor
    _cargos(db, c, dict(B, secuencia=2, porcentaje_secuencia=Decimal(50), porcentaje_retiro=Decimal(50)),
            dict(A, secuencia=1, porcentaje_secuencia=Decimal(50)))
    h1, _ = _haber_mensual(client, admin, db, c)
    _cargos(db, c, dict(B, secuencia=2, porcentaje_secuencia=Decimal(50), porcentaje_retiro=Decimal(50)),
            dict(A, secuencia=1, porcentaje_secuencia=Decimal(50), porcentaje_retiro=Decimal(50)))
    h2, _ = _haber_mensual(client, admin, db, c)
    assert abs(h2 - h1 / 2) < Decimal("0.001")


def test_tres_cargos_y_pension_con_dos_cargos(db, client, admin):
    c = _causante(db)
    _cargos(db, c, dict(A, porcentaje_secuencia=Decimal("33.3333")), dict(B, porcentaje_secuencia=Decimal("33.3333")),
            dict(B, clase=12, porcentaje_secuencia=Decimal("33.3334")))
    haber, rec = _haber_mensual(client, admin, db, c)
    assert len(rec["cargos"]) == 3 and haber > 0
    db.add(m.Beneficiario(causante_id=c.id, apellido="P", nombre="Ana", porcentaje=Decimal(100)))
    db.commit()
    pension, rec_p = _haber_mensual(client, admin, db, c, tipo="pension")
    assert abs(pension - haber * Decimal("0.75")) < Decimal("0.001")
    assert len(rec_p["cargos"]) == 3


@pytest.mark.parametrize("cargos, texto", [
    ([dict(A), dict(B)], "suman 200"),                                                    # ambos con el 100% por defecto
    ([dict(A, porcentaje_secuencia=Decimal(60)), dict(B, porcentaje_secuencia=Decimal(30))], "suman 90"),
    ([dict(A, secuencia=1, porcentaje_secuencia=Decimal(50)), dict(B, secuencia=1, porcentaje_secuencia=Decimal(50))], "mismo número de secuencia"),
])
def test_validacion_de_secuencias(db, client, admin, cargos, texto):
    c = _causante(db)
    _cargos(db, c, *cargos)
    liq = _liq(db, c)
    r = client.post(f"/api/v1/liquidaciones/{liq.id}/calcular", headers=admin)
    assert r.status_code == 422 and texto in r.text


def test_planillas_con_dos_cargos_muestran_el_haber_ponderado(db, client, admin):
    import io
    import openpyxl
    c = _causante(db)
    _cargos(db, c, dict(A, porcentaje_secuencia=Decimal(60)), dict(B, porcentaje_secuencia=Decimal(40)))
    db.add(m.Beneficiario(causante_id=c.id, apellido="P", nombre="Ana", porcentaje=Decimal(100)))
    db.commit()
    liq = _liq(db, c, tipo="pension")
    rec = client.post(f"/api/v1/liquidaciones/{liq.id}/calcular", headers=admin).json()["recibos"][0]
    x = client.get(f"/api/v1/liquidaciones/{liq.id}/planilla.xlsx", headers=admin)
    det = openpyxl.load_workbook(io.BytesIO(x.content))["Haber 1"]
    filas = {det[f"A{r}"].value: det[f"F{r}"].value for r in range(4, det.max_row + 1) if det[f"A{r}"].value}
    assert abs(filas["Total ponderado"] - float(rec["haber_ponderado"])) < 0.005
    assert {det[f"A{r}"].value for r in range(4, det.max_row + 1)} >= {1, 2, "Beneficio"}
    p = client.get(f"/api/v1/liquidaciones/{liq.id}/planilla.pdf", headers=admin)
    assert p.status_code == 200 and p.content.startswith(b"%PDF")
