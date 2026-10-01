"""Importa las liquidaciones proyectadas de renuncia/baja condicionada (una hoja por caso) a JSON.

Fuente: `liquidaciones_complejas_OTP_retiros.xlsx` (transcripción de los PDF de la oficina). Cada hoja tiene
1. Datos del expediente, 2. Cómputo de servicios, 3. Imputación, 4. Encasillamiento definitivo,
5. Liquidación proyectada y el anexo de zona.

Genera:
  backend/app/db/seed_data/casos_proyectados.json   datos de cada caso (puntos 1-4, destinos de zona, mes base) y los
                                                     valores esperados de la planilla (golden de tests/test_golden_proyectadas.py)
  backend/app/db/seed_data/reglas_proyectadas.json  escalas y parámetros PARCIALES (solo lo que aparece en los casos) de las
                                                     vigencias que no cubre la planilla de retroactivos: penitenciaría
                                                     04/2022, 07/2022 y 09/2025; policía 05/2026

Uso:  python scripts/import_casos_proyectados.py <liquidaciones_complejas_OTP_retiros.xlsx>
"""
import json
import re
import sys
from collections import defaultdict
from datetime import date
from decimal import Decimal
from pathlib import Path

import openpyxl

RAIZ = Path(__file__).resolve().parent.parent
SALIDA = RAIZ / "backend" / "app" / "db" / "seed_data"
RETROACTIVOS = SALIDA / "reglas_planillas.json"

MESES = {"enero": 1, "febrero": 2, "marzo": 3, "abril": 4, "mayo": 5, "junio": 6, "julio": 7, "agosto": 8,
         "septiembre": 9, "setiembre": 9, "octubre": 10, "noviembre": 11, "diciembre": 12}

# Clase de los cargos que aparecen como "cargo base" de zona. Penitenciaría: inferidas de la escala (supuesto a confirmar).
CLASE_DE_GRADO = {
    "policia": {"OFICIAL INSPECTOR": 14},
    "penitenciario": {
        "AGENTE": 1, "SUB-OFICIAL SUB-AYUDANTE": 2, "SUB-OFICIAL AYUDANTE": 3, "OFICIAL ADJUTOR": 9,
        "OFICIAL ADJUTOR PRINCIPAL": 10, "SUB-ALCAIDE": 11, "OFICIAL SUB-ALCAIDE": 11, "OFICIAL ALCAIDE": 12,
        "OFICIAL ALCAIDE MAYOR": 13, "OFICIAL PREFECTO": 14,
    },
}

COMPUTO = {
    "periodo cuadro de servicio": "cuadro_servicio",
    "hasta renuncia condicionada": "hasta_renuncia",
    "s. adicionales": "servicios_adicionales",
    "beneficio por título": "beneficio_titulo",
    "suspensiones": "suspensiones",
}
COMPUTO_TOTALES = {"total servicios": "total_servicios", "subtotal": "subtotal", "corresponde": "corresponde"}

# Códigos del encasillamiento cuyo valor es una cantidad (no un porcentaje)
CANTIDADES = {"10", "80"}
# Montos fijos por escalafón (parámetro MONTO_xx o MONTO_xx_PEN)
MONTOS = ("50", "64", "65", "66", "73", "74", "75")


def fecha(s):
    if s in (None, "", "-"):
        return None
    if isinstance(s, date):
        return s.isoformat()
    d, m, a = (int(x) for x in str(s).strip().split("/"))
    return date(a, m, d).isoformat()


def num(x):
    return None if x is None or isinstance(x, str) else Decimal(str(x))


def pct(x):
    """Fracción de la planilla (0,66) → porcentaje 0-100 (66)."""
    return None if x is None else (x * 100).normalize()


def s(x):
    return None if x is None else str(x).strip()


def mes_base(texto):
    t = texto.lower()
    for nombre, n in MESES.items():
        m = re.search(rf"{nombre}\s+(\d{{4}})", t)
        if m:
            return f"{m.group(1)}-{n:02d}"
    raise ValueError(f"No se reconoce el mes base en {texto!r}")


def leer_hoja(ws):
    filas = [list(r) + [None] * 10 for r in ws.iter_rows(values_only=True)]
    seccion = None
    caso = {"hoja": ws.title, "expediente": {}, "computo": [], "computo_esperado": {}, "imputacion": {},
            "encasillamiento": [], "liquidacion": [], "zonas": [], "zona_esperada": {}, "notas": []}
    cab_zona = None
    for r in filas:
        a, b, c, d, e = r[0], r[1], r[2], r[3], r[4]
        if isinstance(a, str):
            t = a.strip().upper()
            if re.match(r"^[1-5]\. ", t):
                seccion = t[0]
                continue
            if t.startswith("ANEXO"):
                seccion = "zona"
                continue
            if t.startswith("VISADO") or t.startswith("NOTAS"):
                seccion = "notas"
                continue
        if seccion == "1" and b:
            caso["expediente"][s(b)] = s(c)
        elif seccion == "2" and b and b != "Concepto":
            k = s(b).lower()
            amd = [int(x or 0) for x in (c, d, e)]
            if k in COMPUTO:
                if any(amd):
                    caso["computo"].append({"concepto": COMPUTO[k], "anios": amd[0], "meses": amd[1], "dias": abs(amd[2])
                                            if COMPUTO[k] == "suspensiones" else amd[2]})
            elif k.startswith("antigüedad final"):
                caso["computo_esperado"]["antiguedad_final"] = int(c)
            else:
                for pref, clave in COMPUTO_TOTALES.items():
                    if k.startswith(pref):
                        caso["computo_esperado"][clave] = amd
        elif seccion == "3" and b and b != "Campo":
            caso["imputacion"][s(b)] = s(c)
        elif seccion == "4" and isinstance(a, str) and a.strip().isdigit():
            caso["encasillamiento"].append({"codigo": str(int(a)), "descripcion": s(b), "valor": c})
        elif seccion == "5":
            if b == "Porcentaje de retiro":
                caso["porcentaje_retiro"] = num(c)
            elif isinstance(a, int):
                caso["liquidacion"].append({"codigo": str(a), "descripcion": s(b), "base": num(c), "unidad": num(d),
                                            "importe": num(e)})
            elif isinstance(b, str) and ("TOTAL DE HABERES" in b or "HABER BRUTO" in b):
                caso["total_haberes"] = num(e)
            elif isinstance(b, str) and b.startswith("PORCENTAJE RETIRO"):
                caso["haber_retiro"] = num(e)
            elif b == "Vigencia":
                caso["vigencia"] = s(c)
            elif isinstance(b, str) and b.startswith("*"):
                caso["notas"].append(s(b))
        elif seccion == "zona":
            if a == "Desde":
                cab_zona = [s(x) for x in r[:10]]
            elif b == "Asunto":
                caso["zona_esperada"]["asunto"] = s(c)
            elif b in ("Total", "Suma de períodos"):
                caso["zona_esperada"]["total"] = num(c)
            elif b == "Porcentaje aplicado":
                caso["zona_esperada"]["aplicado"] = num(c)
            elif b in ("Cargo base", "Jerarquía / clase para zona"):
                caso["zona_esperada"]["cargo_base"] = s(c)
            elif b in ("Base (clase)", "Base de zona"):
                caso["zona_esperada"]["base"] = num(c)
            elif cab_zona and isinstance(a, str) and re.match(r"\d\d/\d\d/\d{4}", a):
                fila = dict(zip(cab_zona, r[:10]))
                zona = fila.get("Zona", fila.get("% Zona", fila.get("% de Zona")))
                caso["zonas"].append({"dependencia": s(fila["Dependencia"]), "porcentaje_zona": pct(num(zona)),
                                      "fecha_desde": fecha(fila["Desde"]), "fecha_hasta": fecha(fila["Hasta"]),
                                      "esperado": num(fila.get("Porc-Final", fila.get("Resultado")))})
        elif seccion == "notas" and b and c:
            caso["notas"].append(f"{s(b)}: {s(c)}")
    if not caso["zona_esperada"].get("asunto"):
        caso["zona_esperada"].pop("asunto", None)
    return caso


def armar(caso, escalas_retro):
    exp, imp = caso["expediente"], caso["imputacion"]
    titulo = s(caso["hoja"])
    nombre_completo = exp["Nombre"].replace(",", "").split()
    n_apellido = len(titulo.split())
    rs = imp.get("R.S.")
    escalafon = "penitenciario" if rs == "37" else "policia"
    tramo = (imp.get("T.") or imp.get("Tramo") or "").zfill(2)
    tipo_personal = "superior" if tramo == "02" else "subalterno"
    dni_cuil = exp.get("DNI") or exp.get("DNI / CUIL")
    dni, _, cuil = dni_cuil.partition("/")
    periodo = mes_base(caso["vigencia"])

    enc = []
    clase = anios = zona_enc = None
    for i, it in enumerate(caso["encasillamiento"]):
        v = it["valor"]
        aplica, valor = True, None
        if v == "-":
            aplica = False
        elif v == "$" or v is None:
            valor = None
            aplica = v == "$"
        else:
            valor = Decimal(str(v)) if it["codigo"] in CANTIDADES else pct(Decimal(str(v)))
        if it["codigo"] == "10":
            clase = int(v)
        elif it["codigo"] == "80":
            anios = valor
        elif it["codigo"] == "83":
            zona_enc = valor
        if v is None:
            continue   # celda vacía (p. ej. resp. jerárquica sin valor): no se encasilla
        enc.append({"codigo": it["codigo"], "descripcion": it["descripcion"], "valor": valor, "aplica": aplica, "orden": i})

    ze = caso["zona_esperada"]
    zona_clase, cargo_base = 0, ze.get("cargo_base")
    if cargo_base:
        grado = cargo_base.split(" - ")[0].strip().upper()
        zona_clase = CLASE_DE_GRADO[escalafon].get(grado) or CLASE_DE_GRADO[escalafon].get("OFICIAL " + grado)
        if zona_clase is None:
            raise ValueError(f"{caso['hoja']}: no hay clase para el cargo base de zona {cargo_base!r}")

    asunto = (caso["hoja_titulo"].split("—")[0].strip()) if caso.get("hoja_titulo") else "RENUNCIA CONDICIONADA"
    cargo_txt = exp.get("Cargo")
    return {
        "hoja": caso["hoja"],
        "escalafon": escalafon,
        "periodo": periodo,
        "asunto": asunto,
        "causante": {
            "expediente": exp.get("Expediente"), "dni": re.sub(r"\D", "", dni), "cuil": re.sub(r"\D", "", cuil) or None,
            "apellido": " ".join(nombre_completo[:n_apellido]), "nombre": " ".join(nombre_completo[n_apellido:]),
            "escalafon": escalafon, "tipo_personal": tipo_personal, "grado": cargo_txt,
            "cuerpo": exp.get("Cuerpo"), "condicion": exp.get("Condición"), "familia": exp.get("Familia"),
            "fecha_ingreso": fecha(exp.get("Ingreso")), "fecha_renuncia_condicionada": fecha(exp.get("Renuncia Condicionada")),
            "cuadro_servicio_desde": fecha(exp.get("Cuadro de Servicio desde")),
            "cuadro_servicio_hasta": fecha(exp.get("Cuadro de Servicio hasta")),
        },
        "computo": caso["computo"],
        "cargo": {
            "secuencia": 1, "porcentaje_secuencia": 100, "clase": clase, "grado": cargo_txt,
            "anios_antiguedad": anios or 0, "zona_porcentaje": zona_enc or 0, "zona_clase": zona_clase,
            "zona_cargo_base": cargo_base,
            "caracter": imp.get("Carácter"), "jurisdiccion": imp.get("Jurisdicción"),
            "unidad_organizativa": imp.get("Unidad Organizativa"), "finalidad": imp.get("Finalidad"),
            "funcion": imp.get("Función"), "regimen_salarial": rs, "agrupamiento": imp.get("A."), "tramo": tramo,
            "subtramo": imp.get("ST.") or imp.get("Subtramo"),
        },
        "encasillamiento": enc,
        "zonas": [{k: v for k, v in z.items() if k != "esperado"} for z in caso["zonas"]],
        "esperado": {
            "computo": caso["computo_esperado"],
            "porcentaje_retiro": pct(caso["porcentaje_retiro"]),
            "conceptos": {l["codigo"]: l["importe"] for l in caso["liquidacion"]},
            "bases": {l["codigo"]: l["base"] for l in caso["liquidacion"] if l["base"] is not None},
            "total_haberes": caso["total_haberes"],
            "haber_retiro": caso["haber_retiro"],
            "zona": {"filas": [z["esperado"] for z in caso["zonas"]], "total": ze.get("total"),
                     "aplicado": pct(ze["aplicado"]) if ze.get("aplicado") is not None else pct(zona_enc / 100 if zona_enc else Decimal(0)),
                     "base": ze.get("base")},
        },
        "notas": caso["notas"],
    }


def reglas_parciales(casos, escalas_retro):
    """Escalas y parámetros por (escalafón, mes) deducidos de las liquidaciones; reporta conflictos."""
    vig = defaultdict(lambda: {"escala": {}, "parametros": {}})
    discrepancias = []

    def poner(dic, clave, valor, caso, que):
        if valor is None:
            return
        if clave in dic and dic[clave] != valor:
            discrepancias.append(f"{caso['hoja']}: {que} {clave} = {valor} (otro caso: {dic[clave]})")
            return
        dic[clave] = valor

    for c in casos:
        esc, per = c["escalafon"], c["periodo"]
        retro = escalas_retro.get(per) if esc == "policia" else None
        imp = c["esperado"]["conceptos"]
        bases = c["esperado"]["bases"]
        sfx = "_PEN" if esc == "penitenciario" else ""
        clases = {c["cargo"]["clase"]: imp["10"]}
        if c["cargo"]["zona_clase"] and c["esperado"]["zona"]["base"]:
            clases[c["cargo"]["zona_clase"]] = c["esperado"]["zona"]["base"]
        if retro is not None:   # policía con escala ya cargada desde la planilla de retroactivos: solo se controla
            for k, v in clases.items():
                if Decimal(str(retro.get(str(k)))) != v:
                    discrepancias.append(f"{c['hoja']}: clase {k} = {v} y la planilla de retroactivos tiene {retro.get(str(k))}")
            continue
        v = vig[(esc, per)]
        for k, val in clases.items():
            poner(v["escala"], k, val, c, "clase")
        p = v["parametros"]
        for cod in ("23", "58", "51"):
            if bases.get(cod):
                poner(p, "BASE_JEFE" + sfx, bases[cod], c, "base")
        if bases.get("24"):
            poner(p, ("BASE_TITULO_PEN" if sfx else "BASE_JEFE"), bases["24"], c, "base título")
        for cod in MONTOS:
            if imp.get(cod):
                poner(p, f"MONTO_{cod}{sfx}", imp[cod], c, "monto")
        if sfx:
            enc = {e["codigo"]: e["valor"] for e in c["encasillamiento"] if e["valor"] is not None}
            for cod, campo in (("90", "PORC_ADIC_FZA_SEG"), ("91", "PORC_AUM_0705"), ("92", "PORC_AUM_0907"),
                               ("93", "PORC_AUM_0308"), ("51", "PORC_PERS_SEG_2022"), ("52", "PORC_FORTALECIMIENTO")):
                if cod in enc:
                    poner(p, campo + sfx, enc[cod], c, "porcentaje")
            if imp.get("51") and bases.get("51") is None and p.get("BASE_JEFE_PEN"):   # "$": % = importe / base
                poner(p, "PORC_PERS_SEG_2022_PEN", (imp["51"] / p["BASE_JEFE_PEN"] * 100).quantize(Decimal("1")), c, "porcentaje")
            poner(p, "PORC_ANTIGUEDAD_PEN", Decimal(2), c, "porcentaje")
    # El 51 penitenciario con "$" (sin base visible) en 07/2022
    for (esc, per), v in vig.items():
        if esc == "penitenciario" and "PORC_PERS_SEG_2022_PEN" not in v["parametros"]:
            for c in casos:
                if c["escalafon"] == esc and c["periodo"] == per and c["esperado"]["conceptos"].get("51"):
                    v["parametros"]["PORC_PERS_SEG_2022_PEN"] = (c["esperado"]["conceptos"]["51"]
                                                                 / v["parametros"]["BASE_JEFE_PEN"] * 100).quantize(Decimal("1"))
    out = []
    for (esc, per), v in sorted(vig.items()):
        anio, mes = int(per[:4]), int(per[5:])
        hasta = (date(anio + (mes == 12), mes % 12 + 1, 1).toordinal() - 1)
        out.append({"escalafon": esc, "desde": f"{per}-01", "hasta": date.fromordinal(hasta).isoformat(),
                    "escala": {str(k): v["escala"][k] for k in sorted(v["escala"])}, "parametros": v["parametros"]})
    return out, discrepancias


def _json(o):
    if isinstance(o, Decimal):
        return int(o) if o == o.to_integral_value() else float(o)
    raise TypeError(type(o))


def main(xlsx):
    wb = openpyxl.load_workbook(xlsx, data_only=True)
    retro = json.loads(RETROACTIVOS.read_text(encoding="utf-8"))
    escalas_retro = {e["desde"][:7]: e["clases"] for e in retro["escalas"]}
    casos = []
    for ws in wb.worksheets:
        crudo = leer_hoja(ws)
        crudo["hoja_titulo"] = ws["A1"].value
        casos.append(armar(crudo, escalas_retro))
    vigencias, discrepancias = reglas_parciales(casos, escalas_retro)
    fuente = Path(xlsx).name
    (SALIDA / "casos_proyectados.json").write_text(json.dumps(
        {"fuente": fuente, "casos": casos}, default=_json, ensure_ascii=False, indent=1), encoding="utf-8")
    (SALIDA / "reglas_proyectadas.json").write_text(json.dumps(
        {"fuente": fuente, "provisorio": True, "vigencias": vigencias, "discrepancias": discrepancias},
        default=_json, ensure_ascii=False, indent=1), encoding="utf-8")
    print(f"{len(casos)} casos, {len(vigencias)} vigencias parciales")
    for d in discrepancias:
        print("  discrepancia:", d)


if __name__ == "__main__":
    if len(sys.argv) != 2:
        sys.exit(__doc__)
    main(sys.argv[1])
