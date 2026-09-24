#!/usr/bin/env python3
"""Genera los casos dorados (valores esperados) recalculando la planilla original con LibreOffice.

Carga los datos de cada caso en la hoja DATOS de "BASE DE LIQ RETIROS-REAJUSTES-PENSION POLICIA", deja en las
planillas de liquidación solo los tramos de jul/2022 a mar/2025 (que incluyen 5 SAC), recalcula con
LibreOffice y guarda los resultados en backend/tests/golden/casos_planilla.json. Los valores esperados NO salen del
motor propio: son los de las fórmulas de Excel.

Uso:  python scripts/golden_planilla.py <BASE DE LIQ ....xlsx convertido>
Requiere openpyxl y soffice (LibreOffice) en el PATH.
"""
import json
import re
import shutil
import subprocess
import sys
import tempfile
from datetime import date, datetime, timedelta
from pathlib import Path

import openpyxl
from openpyxl.cell.cell import MergedCell

OUT = Path(__file__).resolve().parent.parent / "backend/tests/golden/casos_planilla.json"
# Diseño de cada hoja de liquidación: filas de tramos jul/2022..mar/2025 (5 SAC dentro), filas a vaciar,
# celdas de subtotal / descuentos / líquido y celdas de anticipo y haberes devengados a poner en 0.
# B2..B5 están corridas una fila respecto de Retiro/B1. Además `B2!D122` calcula el OSEP sobre `B3!F118`
# (error de la planilla): en B2..B5 los descuentos y el líquido NO se toman de Excel sino que se derivan (6%).
DISENO = {
    "Planilla Retiro": dict(filas=range(75, 112), limpiar=list(range(14, 75)) + list(range(112, 117)),
                            subtotal="F117", desc="D125", liquido="D131", ceros=["D127"], anticipo="D128", excel_desc=True),
    "B1 ": dict(filas=range(75, 112), limpiar=list(range(14, 75)) + list(range(112, 117)),
                subtotal="F117", desc="D124", liquido="D131", ceros=["D126", "D130"], anticipo="D127", excel_desc=True),
}
for _h in ("B2", "B3", "B4", "B5"):
    DISENO[_h] = dict(filas=range(76, 113), limpiar=list(range(14, 76)) + list(range(113, 118)),
                      subtotal="F118", desc="D125", liquido=None, ceros=["D127", "D131"], anticipo="D128", excel_desc=False)
RANGO = {"desde": "2022-07-01", "hasta": "2025-03-31"}

# Filas de DATOS por beneficiario: (porcentaje, art. 37). B1 no tiene art. 37 en la planilla.
BENEF_DATOS = [("D6", None), ("D11", "C15"), ("D17", "C21"), ("D23", "C27"), ("D29", "C33"), ("D35", "C39")]
HOJA_BENEF = ["B1 ", "B2", "B3", "B4", "B5"]

CASOS = [
    {
        "nombre": "retiro_superior_clase16_todos_los_adicionales",
        "tipo": "retiro",
        "causante": {"tipo_personal": "superior"},
        "cargo": {"clase": 16, "responsabilidad_jerarquica_porcentaje": 81.01, "recargo_servicio_porcentaje": 66,
                  "titulo": "grado", "titulo_pregrado_nivel": 0, "riesgo_especial": True, "zona_porcentaje": 10,
                  "zona_clase": 16, "anios_antiguedad": 30, "cuerpo_apoyo_porcentaje": 10, "adicional_seguridad": True,
                  "porcentaje_retiro": 100},
        "anticipo": 150000,
    },
    {
        "nombre": "retiro_subalterno_clase9_pregrado",
        "tipo": "retiro",
        "causante": {"tipo_personal": "subalterno"},
        "cargo": {"clase": 9, "titulo": "ninguno", "titulo_pregrado_nivel": 2, "anios_antiguedad": 25,
                  "porcentaje_retiro": 100},
        "anticipo": 0,
    },
    {
        "nombre": "retiro_clase20_porcentaje_parcial",
        "tipo": "retiro",
        "causante": {"tipo_personal": "superior"},
        "cargo": {"clase": 20, "responsabilidad_jerarquica_porcentaje": 100, "titulo": "posgrado", "riesgo_especial": True,
                  "anios_antiguedad": 27, "adicional_seguridad": False, "porcentaje_retiro": 87},
        "anticipo": 0,
    },
    {
        "nombre": "pension_conyuge_y_dos_hijos_con_art37",
        "tipo": "pension",
        "causante": {"tipo_personal": "superior"},
        "cargo": {"clase": 14, "titulo": "ninguno", "titulo_pregrado_nivel": 1, "zona_porcentaje": 5, "zona_clase": 14,
                  "anios_antiguedad": 28, "adicional_seguridad": True, "cuerpo_apoyo_porcentaje": 0, "porcentaje_retiro": 96},
        "beneficiarios": [
            {"parentesco": "conyuge", "porcentaje": 50, "art37": False},
            {"parentesco": "hijo", "porcentaje": 25, "art37": True},
            {"parentesco": "hijo", "porcentaje": 25, "art37": True},
        ],
        "anticipo": 0,
        # `B3` de la planilla apunta a las columnas equivocadas en ene/feb 2023 ('2023'!O70 y T70 en vez de E70 y J70),
        # por eso solo se toman B1 y B2 y el tercer beneficiario (igual al segundo) se compara con B2.
        "hojas": ["B1 ", "B2"],
        "esperado_por_recibo": [0, 1, 1],
    },
]


def cargar_datos(wb, caso):
    ds = wb["DATOS"]
    c = caso["cargo"]
    ds["D47"] = c["clase"]
    for celda in ("D49", "D51"):
        ds[celda] = c.get("responsabilidad_jerarquica_porcentaje", 0) / 100
    ds["D55"] = c.get("recargo_servicio_porcentaje", 0) / 100
    ds["D57"] = {"ninguno": 0, "grado": 0.20, "posgrado": 0.25}[c.get("titulo", "ninguno")]
    ds["D59"] = c.get("titulo_pregrado_nivel", 0)
    riesgo = c.get("riesgo_especial", False)
    ds["D61"] = 0.065 if riesgo else 0
    ds["D63"] = 0.13 if riesgo else 0
    ds["E65"] = c.get("zona_clase", 0)
    ds["D65"] = c.get("zona_porcentaje", 0) / 100
    ds["D67"] = c.get("anios_antiguedad", 0)
    ds["D71"] = c.get("cuerpo_apoyo_porcentaje", 0) / 100
    ds["D73"] = 1 if c.get("adicional_seguridad") else 0
    ds["D75"] = 0
    ds["D77"] = c.get("porcentaje_retiro", 0) / 100
    benef = caso.get("beneficiarios", [])
    for i, (celda_pct, celda_art) in enumerate(BENEF_DATOS):
        b = benef[i] if i < len(benef) else None
        ds[celda_pct] = (b["porcentaje"] / 100) if b else 0
        if celda_art:
            ds[celda_art] = 0.05 if (b and b["art37"]) else 0


DUPLICADA = re.compile(r"^=\+?(.+)\+\1$")


def corregir_duplicados(ws, hoja):
    """Errores de tipeo de la planilla: `=+'2023'!J68+'2023'!J68` suma dos veces la misma celda. Se corrigen en la
    copia de trabajo (referencia única) y se informan, para que el valor esperado sea el cálculo previsto."""
    corregidas = []
    for r in DISENO[hoja]["filas"]:
        v = ws[f"A{r}"].value
        m = DUPLICADA.match(v) if isinstance(v, str) else None
        if m:
            ws[f"A{r}"] = "=+" + m.group(1)
            corregidas.append({"hoja": hoja, "fila": r, "original": v, "corregida": "=+" + m.group(1)})
    return corregidas


def preparar_hoja(ws, caso, hoja):
    cfg = DISENO[hoja]
    for r in cfg["limpiar"]:
        for idx in range(1, 7):
            celda = ws.cell(r, idx)
            if not isinstance(celda, MergedCell):
                celda.value = None
    for celda in cfg["ceros"]:
        ws[celda] = 0
    ws[cfg["anticipo"]] = caso.get("anticipo", 0)


def leer(ws, hoja):
    cfg = DISENO[hoja]
    tramos, sacs = [], []
    for r in cfg["filas"]:
        a, b, c, dmes, f = (ws[f"{col}{r}"].value for col in "ABCDF")
        if b == "SAC":
            sacs.append({"fila": r, "importe": f})
        else:
            tramos.append({"fila": r, "desde": b.date().isoformat(), "meses": dmes, "nominal": a, "total": f})
    return {
        "tramos": tramos, "sacs": sacs, "subtotal_credito": ws[cfg["subtotal"]].value,
        "descuentos": ws[cfg["desc"]].value if cfg["excel_desc"] else None,
        "liquido": ws[cfg["liquido"]].value if cfg["liquido"] else None,
        "descuentos_de_excel": cfg["excel_desc"],
    }


def main(src):
    src = Path(src)
    resultado = {"rango": RANGO, "fuente": src.name, "casos": []}
    with tempfile.TemporaryDirectory() as tmp:
        tmp = Path(tmp)
        for caso in CASOS:
            wb = openpyxl.load_workbook(src)
            cargar_datos(wb, caso)
            es_retiro = caso["tipo"] == "retiro"
            hojas = ["Planilla Retiro"] if es_retiro else caso["hojas"]
            correcciones = []
            for h in hojas:
                correcciones += corregir_duplicados(wb[h], h)
                preparar_hoja(wb[h], caso, h)
            f_in = tmp / f"{caso['nombre']}.xlsx"
            wb.save(f_in)
            out_dir = tmp / "out"
            out_dir.mkdir(exist_ok=True)
            subprocess.run(["soffice", "--headless", "--convert-to", "xlsx", "--outdir", str(out_dir), str(f_in)],
                           check=True, capture_output=True, timeout=300)
            calc = openpyxl.load_workbook(out_dir / f_in.name, data_only=True)
            esperado = [dict(hoja=h, **leer(calc[h], h)) for h in hojas]
            for e in esperado:
                assert isinstance(e["subtotal_credito"], (int, float)), (caso["nombre"], e["hoja"], e["subtotal_credito"])
                assert e["liquido"] is None or isinstance(e["liquido"], (int, float)), (caso["nombre"], e["hoja"], e["liquido"])
            resultado["casos"].append({k: v for k, v in caso.items()} | {"esperado": esperado, "correcciones_planilla": correcciones})
            for c in correcciones:
                print("   corrección:", c["hoja"], "fila", c["fila"], c["original"], "->", c["corregida"])
            print(caso["nombre"], [(e["hoja"], round(e["subtotal_credito"], 2), e["liquido"] and round(e["liquido"], 2)) for e in esperado])
    OUT.write_text(json.dumps(resultado, ensure_ascii=False, indent=1, default=str), encoding="utf-8")
    print("->", OUT)


if __name__ == "__main__":
    if len(sys.argv) != 2:
        sys.exit(__doc__)
    main(sys.argv[1])
